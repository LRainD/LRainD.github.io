from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
BUILD = ROOT / "analysis" / "docx_build"
EQUATIONS = BUILD / "equations"
OUTPUT_DIR = ROOT / "outputs" / "20260916_集采计划预测"
OUTPUT = OUTPUT_DIR / "集采计划预测模型算法与指标说明.docx"

NAVY = "0B355C"
BLUE = "17679F"
ORANGE = "ED7D31"
INK = "182B3D"
MUTED = "687B8E"
LINE = "D9D9D9"
PALE = "F3F7FA"
BODY_FONT = "Heiti SC"


def set_run_font(run, name=BODY_FONT, size=10.5, bold=False, color=INK):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia"):
        fonts.set(qn(f"w:{attr}"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color=LINE, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        border = borders.find(qn(tag))
        if border is None:
            border = OxmlElement(tag)
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), size)
        border.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=105, bottom=100, end=105):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_row_no_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:cantSplit")
    tr_pr.append(node)


def set_keep_with_next(paragraph, value=True):
    ppr = paragraph._p.get_or_add_pPr()
    node = ppr.find(qn("w:keepNext"))
    if node is None:
        node = OxmlElement("w:keepNext")
        ppr.append(node)
    node.set(qn("w:val"), "1" if value else "0")


def remove_paragraph_border(element):
    ppr = element.get_or_add_pPr()
    border = ppr.find(qn("w:pBdr"))
    if border is not None:
        ppr.remove(border)


def style_paragraph(paragraph, size=10.5, color=INK, bold=False, align=None, before=0, after=6, line=1.25):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align
    for run in paragraph.runs:
        set_run_font(run, size=size, bold=bold, color=color)
    return paragraph


def add_body(doc, text, *, bold_lead=None, after=6):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        run = p.add_run(text[len(bold_lead):])
        set_run_font(run)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return style_paragraph(p, after=after)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.24)
        p.paragraph_format.first_line_indent = Inches(-0.14)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.18
        run = p.add_run(item)
        set_run_font(run)


def add_equation(doc, name, width=6.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(7)
    p.add_run().add_picture(str(EQUATIONS / f"{name}.png"), width=Inches(width))
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    set_run_font(run, size=17 if level == 1 else 13, bold=True, color="000000")
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    return p


def add_table(doc, headers, rows, widths, font_size=9.5, alignments=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    set_repeat_header(header)
    for idx, (cell, label, width) in enumerate(zip(header.cells, headers, widths)):
        cell.width = Inches(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(label)
        set_run_font(run, size=9.5, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_no_split(row)
        for idx, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE)
            p = cell.paragraphs[0]
            p.alignment = (alignments[idx] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            run = p.add_run(str(value))
            set_run_font(run, size=font_size, color=INK)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    end = paragraph.add_run(" 页")
    set_run_font(end, size=8.5, color=MUTED)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for level in (1, 2):
        style = styles[f"Heading {level}"]
        style.font.name = BODY_FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    title_style = styles["Title"]
    title_style.font.name = BODY_FONT
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    remove_paragraph_border(title_style._element)

    # Cover
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(82)
    r = p.add_run("中建电商智慧供采事业部 AI 创新大赛")
    set_run_font(r, size=10.5, bold=True, color=ORANGE)
    title = doc.add_paragraph(style="Title")
    remove_paragraph_border(title._p)
    title.paragraph_format.space_before = Pt(14)
    title.paragraph_format.space_after = Pt(12)
    run = title.add_run("集采计划预测模型算法与指标说明")
    set_run_font(run, size=25, bold=True, color="000000")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run("面向演示讲解和评委问答的技术说明")
    set_run_font(run, size=13, color=MUTED)
    add_body(doc, "本说明对应当前演示系统和预测代码，解释候选分、项目暴露度、评价指标、四类预测算法、加权融合、20类特征和金额相似度。重点结论是：候选分用于候选排序和置信标识，预测效果则由2026年真实计划的一对一匹配结果评价，两者不能混为一谈。", after=14)
    add_body(doc, "数据范围  中国建筑一局（集团）有限公司2023年至2026年集采计划、历史招标采购数据、项目数据及基础字典", after=5)
    add_body(doc, "编制日期  2026年9月", after=5)
    doc.add_page_break()

    add_heading(doc, "阅读要点", 1)
    add_bullets(doc, [
        "去年延续基线的70%是人为设定的固定候选分，不是经过概率校准的70%发生概率。",
        "项目暴露度只反映各组织在当前项目表中的项目数量，尚未纳入项目状态、合同额和建设阶段。",
        "精确率衡量预测结果中有多少命中，召回率衡量真实计划中有多少被提前覆盖。",
        "融合模型权重由人工设定，当前代码没有保存自动寻优过程，答辩时不应表述为模型自动学习所得。",
        "单条金额相似度使用对数距离，总体金额准确度使用SMAPE，两者是不同指标。",
    ])
    add_heading(doc, "候选实体的含义", 2)
    add_body(doc, "系统先把组织机构、区域、采购品类和采购内容拼成一个候选实体。例如，一局三公司、华南、专业分包、防水专业分包构成一个独立候选。模型判断的是该组合在目标年度是否可能再次成为集采计划条目。")

    add_heading(doc, "1 去年延续基线候选分", 1)
    add_body(doc, "去年延续基线把2025年所有非租赁计划条目平移到2026年，并为每条记录统一赋值0.70。该分数是基线模型的排序分和置信标识，不是根据历史样本训练得到的概率。")
    add_table(doc, ["问题", "当前实现"], [
        ("是否决定进入预测清单", "不决定。2025年非租赁条目全部延续。"),
        ("是否用于候选排序", "基线条目全部同分，实际没有区分排序。"),
        ("是否影响置信度", "影响。0.70在当前阈值下归为中置信度。"),
        ("是否参与真实结果匹配", "不参与。匹配使用内容、组织、区域、类型、方式、时间和金额。"),
    ], [2.05, 4.85], font_size=9.7)
    add_body(doc, "答辩建议表述  70%是基线模型统一设置的候选分，用于展示置信程度。它不是严格统计意义上的发生概率，也不决定基线条目是否入选。", bold_lead="答辩建议表述  ")

    add_heading(doc, "2 项目暴露度", 1)
    add_body(doc, "项目暴露度用于表达一个组织产生采购需求的潜在基础。代码按项目数据中的所属组织统计项目记录数，再使用对数归一化将其压缩到0至1。")
    add_equation(doc, "project_exposure", 3.2)
    add_body(doc, "其中，nₒ是组织o的项目记录数，nₘₐₓ是项目数最多组织的记录数。项目最多的组织取值为1；没有匹配到项目的组织取值为0。对数压缩可以降低超大组织对结果的支配程度。")
    add_table(doc, ["模型或环节", "项目暴露度的作用"], [
        ("指数加权", "直接权重2%"),
        ("阻尼趋势", "以5%的系数加入趋势原始分，再做排名归一化"),
        ("随机森林和梯度提升", "作为第20个模型特征"),
        ("加权融合", "直接权重4%"),
    ], [2.2, 4.7], font_size=9.7)
    add_body(doc, "当前限制  项目数据只有6个字段，代码没有按项目状态、合同额、开竣工时间或区域筛选。因此，项目暴露度表示项目记录规模，不等于当前在施项目规模。", bold_lead="当前限制  ")

    add_heading(doc, "3 精确率召回率和综合相似度", 1)
    add_heading(doc, "一对一匹配", 2)
    add_body(doc, "系统先在相同采购品类中寻找预测条目和实际条目的候选组合，并计算单条匹配得分。匹配得分达到58%后才有资格成为命中。系统按得分从高到低选择，保证一条预测最多匹配一条实际记录，一条实际记录也最多使用一次。")
    add_equation(doc, "pair_score", 6.6)
    add_heading(doc, "精确率召回率和F1", 2)
    add_equation(doc, "precision_recall", 5.7)
    add_body(doc, "M是成功匹配数，Nₚ是预测条目数，Nₐ是实际条目数。精确率回答预测出来的条目有多少命中；召回率回答实际发生的条目有多少被提前覆盖；F1用于平衡两者。")
    add_table(doc, ["指标", "去年延续基线", "计算"], [
        ("精确率", "55.77%", "826 ÷ 1,481"),
        ("召回率", "52.68%", "826 ÷ 1,568"),
        ("F1", "54.18%", "精确率与召回率的调和平均"),
    ], [1.45, 1.6, 3.85], font_size=9.7, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_heading(doc, "综合相似度", 2)
    add_equation(doc, "composite", 5.8)
    add_table(doc, ["组成", "权重", "定义"], [
        ("条目F1", "50%", "综合精确率与召回率"),
        ("字段相似度", "20%", "成功匹配条目的内容、组织、区域、采购类型和采购方式平均相似度"),
        ("金额准确度", "15%", "1减去匹配条目的平均SMAPE，最低为0"),
        ("时间准确度", "10%", "计划开始日期相差180天时降为0"),
        ("规模准确度", "5%", "预测总条目数与实际总条目数的接近程度"),
    ], [1.45, 0.75, 4.7], font_size=9.4, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_equation(doc, "accuracy_parts", 6.15)

    add_heading(doc, "4 指数加权模型", 1)
    add_body(doc, "指数加权模型把最近年份的历史延续性放在首位，并用需求热度、招采行为和项目规模进行小幅修正。候选条目评分由八项证据组成。")
    add_equation(doc, "ewma_score", 6.35)
    add_table(doc, ["证据", "权重", "解释"], [
        ("2025年是否出现", "46%", "最近一年出现是最强证据"),
        ("2024年是否出现", "20%", "较近历史延续性"),
        ("2023年是否出现", "9%", "更早历史延续性"),
        ("2025年品类需求热度", "8%", "相同采购品类和采购内容的全局热度"),
        ("2025年组织需求热度", "6%", "相同组织和采购品类的计划热度"),
        ("组织品类招采活跃度", "5%", "相同组织和采购品类的招标采购记录"),
        ("相同候选招采活跃度", "4%", "相同候选实体的招标采购记录"),
        ("项目暴露度", "2%", "组织当前项目数量信号"),
    ], [2.35, 0.7, 3.85], font_size=9.2, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_body(doc, "需求热度和招采活跃度经过对数及双曲正切函数压缩，避免少数极大值支配结果。候选按得分降序排列，再按照年度预测条目数截取。")
    add_heading(doc, "年度条目数和金额", 2)
    add_equation(doc, "ewma_count", 5.3)
    add_body(doc, "当前计算得到1,069条，模型从全部候选中选择得分最高的1,069条。金额采用相同的60%、30%、10%年度权重，并按实际有数据的年份重新归一化。当前实现属于固定时间衰减加权，不是递归形式的标准EWMA。")

    add_heading(doc, "5 随机森林模型", 1)
    add_body(doc, "随机森林中的一棵树是一套由历史数据学习出的分支判断规则。每个节点选择一个特征和阈值，把样本分成两组；叶节点输出其中训练标签的平均值，可理解为该局部样本的再次出现比例。")
    add_body(doc, "树1、树2、树3只是不同树的序号，没有预先规定的业务职责。每棵树使用不同的有放回抽样样本，每次分裂只随机考察约4个特征，因此会形成不同的判断路径。")
    add_table(doc, ["参数", "当前值", "含义"], [
        ("树数量", "70", "70套独立判断规则"),
        ("最大深度", "5", "每棵树最多进行5层分裂"),
        ("最小叶节点", "12", "叶节点至少保留12条训练样本"),
        ("候选特征数", "约4", "每次分裂从20类特征中随机选取约√20个"),
        ("训练标签", "0或1", "下一年没有出现为0，出现为1"),
    ], [1.6, 1.0, 4.3], font_size=9.5, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_equation(doc, "rf", 4.25)
    add_body(doc, "每棵树的分裂目标是尽量降低左右子节点的标签方差。最终候选分等于70棵树输出的平均值。")

    add_heading(doc, "6 梯度提升树模型", 1)
    add_body(doc, "梯度提升树采用顺序修正。第一步用训练标签平均值作为基础预测，之后每棵小树学习当前预测与真实标签之间的残差。")
    add_equation(doc, "gbt", 4.9)
    add_table(doc, ["参数", "当前值", "作用"], [
        ("树数量", "55", "连续进行55轮误差修正"),
        ("学习率", "0.06", "每轮只采纳新树输出的6%"),
        ("最大深度", "2", "使用较简单的小树，降低过拟合风险"),
        ("最小叶节点", "15", "每个局部修正规则至少由15条样本支持"),
    ], [1.6, 1.0, 4.3], font_size=9.5, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_body(doc, "如果真实标签为1而当前预测为0.4，残差为+0.6，新树会推动预测上升。如果真实标签为0而当前预测为0.7，残差为-0.7，新树会推动预测下降。学习率0.06使每次调整保持较小幅度。")

    add_heading(doc, "7 加权融合模型", 1)
    add_equation(doc, "ensemble", 6.5)
    add_body(doc, "R表示排名归一化。不同模型的原始分数尺度可能不同，排名归一化把候选的相对位置转换为0至1，再进行加权。")
    add_table(doc, ["分量", "权重", "选择原因"], [
        ("指数加权", "22%", "保留年度延续规律，稳定且容易解释"),
        ("阻尼趋势", "18%", "表达需求变化方向，但避免机械外推"),
        ("随机森林", "25%", "学习复杂的非线性组合"),
        ("梯度提升树", "25%", "连续修正难预测的候选"),
        ("同候选招采活跃度", "6%", "补充近期业务行为信号"),
        ("项目暴露度", "4%", "补充组织潜在需求基础"),
    ], [2.0, 0.75, 4.15], font_size=9.4, alignments=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT])
    add_heading(doc, "阻尼趋势", 2)
    add_equation(doc, "trend", 4.7)
    add_body(doc, "阻尼趋势只延续上一年度变化幅度的25%。全年条目数还被限制在2025年条目数的90%至110%之间，金额趋势限制在上年金额的50%至150%之间，避免短期波动被完整放大。")
    add_heading(doc, "权重来源", 2)
    add_body(doc, "当前22%、18%、25%、25%、6%、4%是人工设定的工程权重。代码中没有保存网格搜索、贝叶斯优化或交叉验证选权重的过程，因此不能表述为AI自动学习所得。现有结构让树模型合计占50%，历史规则和趋势占40%，招采与项目直接信号占10%。")
    add_body(doc, "后续优化建议  使用2024年和2025年滚动回测，在权重和模型参数的候选范围内寻找综合相似度更高且年度表现稳定的组合，并继续把2026年作为未参与调参的盲测数据。", bold_lead="后续优化建议  ")

    add_heading(doc, "8 二十类模型特征", 1)
    add_body(doc, "以下定义均以预测2026年为例。候选实体e由组织、区域、采购品类和采购内容组成。N(y,e)表示候选e在年份y的计划条目数。ln(1+x)用于压缩数量差异。")
    features = [
        (1, "count_lag1", "2025出现次数", "N(2025,e)"),
        (2, "count_lag2", "2024出现次数", "N(2024,e)"),
        (3, "count_lag3", "2023出现次数", "N(2023,e)"),
        (4, "present_lag1", "2025是否出现", "出现为1，否则0"),
        (5, "present_lag2", "2024是否出现", "出现为1，否则0"),
        (6, "present_lag3", "2023是否出现", "出现为1，否则0"),
        (7, "recency", "历史新近度", "2025出现取1；否则2024取1/2；否则2023取1/3；都未出现取0"),
        (8, "item_demand_lag1", "2025品类内容需求", "ln(1+2025年相同采购品类和采购内容的条目数)"),
        (9, "item_demand_lag2", "2024品类内容需求", "ln(1+2024年相同采购品类和采购内容的条目数)"),
        (10, "org_demand_lag1", "2025组织需求", "ln(1+2025年相同组织和采购品类的条目数)"),
        (11, "org_demand_lag2", "2024组织需求", "ln(1+2024年相同组织和采购品类的条目数)"),
        (12, "region_demand_lag1", "2025区域需求", "ln(1+2025年相同区域和采购品类的条目数)"),
        (13, "amount_lag1", "2025金额", "ln(1+同候选2025年金额中位数)，无有效金额取0"),
        (14, "amount_lag2", "2024金额", "ln(1+同候选2024年金额中位数)，无有效金额取0"),
        (15, "amount_change", "金额变化", "amount_lag1减amount_lag2"),
        (16, "start_month", "开始时间位置", "2025开始日距年初天数÷30.5÷12；缺失取0.5"),
        (17, "duration", "采购周期", "2025周期中位天数截断到0至730天后÷365；缺失取0.12"),
        (18, "transaction_lag1", "组织品类招采", "ln(1+2025年相同组织和采购品类的招标采购次数)"),
        (19, "transaction_entity_lag1", "同候选招采", "ln(1+2025年相同候选实体的招标采购次数)"),
        (20, "project_exposure", "项目暴露度", "ln(1+组织项目数)÷ln(1+最大组织项目数)"),
    ]
    add_table(doc, ["序号", "代码字段", "业务定义", "计算方法"], features, [0.48, 1.65, 1.55, 3.22], font_size=8.8, alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT])
    add_body(doc, "页面分组校正  演示页面目前写作历史延续7项、需求热度6项、金额与时间6项、项目暴露1项。按代码更准确的分组是历史延续7项、需求与招采7项、金额与时间5项、项目暴露1项，总数仍为20项。", bold_lead="页面分组校正  ")

    add_heading(doc, "9 金额相似度示例", 1)
    add_body(doc, "单条预测和实际记录进行匹配时，金额相似度采用对数距离。P表示预测金额，A表示实际金额。")
    add_equation(doc, "amount_similarity", 5.4)
    add_body(doc, "当预测金额为800万元、实际金额为900万元时：")
    add_equation(doc, "amount_example", 4.9)
    add_body(doc, "因此页面显示92.7%。对数公式关注相对差异，降低大金额绝对差额的影响。分母ln5表示当金额比例接近5倍时，相似度接近0。公式中的加1用于兼容金额为0的情况。")
    add_heading(doc, "单条金额相似度和总体金额准确度", 2)
    add_equation(doc, "smape", 5.65)
    add_body(doc, "800万元和900万元的SMAPE为100÷850，即11.76%。如果只有这一条匹配记录，总体金额准确度为88.24%。因此92.7%与88.24%并不矛盾：前者用于单条匹配得分，权重7%；后者用于模型总体综合相似度，权重15%。")

    add_heading(doc, "答辩表述注意事项", 1)
    add_table(doc, ["容易引起误解的说法", "建议表述"], [
        ("候选分70%就是有70%概率发生", "70%是基线模型统一设置的候选分和置信标识，不是经过校准的发生概率。"),
        ("融合权重由AI自动算出", "当前融合权重由人工设定，用于平衡历史规则、树模型和业务信号。"),
        ("项目暴露度代表在施项目规模", "当前项目暴露度只反映项目表中各组织的项目记录数。"),
        ("综合相似度就是平均匹配分", "综合相似度由F1、字段、金额、时间和规模五类指标加权得到。"),
        ("树1负责金额，树2负责区域", "每棵树使用不同样本和随机特征学习完整判断规则，没有固定业务分工。"),
    ], [2.65, 4.25], font_size=9.3)

    add_heading(doc, "实现依据", 1)
    add_body(doc, "本文档对应当前集采计划预测流水线、演示系统和2026年盲测结果。公式、权重、阈值和模型参数均按当前实现整理，后续如调整代码，应同步更新本文档。")

    for sec in doc.sections:
        sec.footer_distance = Inches(0.35)
        add_page_number(sec.footer.paragraphs[0])

    doc.core_properties.title = "集采计划预测模型算法与指标说明"
    doc.core_properties.subject = "集采计划预测算法、特征与评价指标说明"
    doc.core_properties.author = "中建电商智慧供采事业部"
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
